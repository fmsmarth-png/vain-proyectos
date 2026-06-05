import requests
from openpyxl import load_workbook
import json

SUPABASE_URL = "https://swjmqtnhdtiwopexbezx.supabase.co"
SUPABASE_KEY = "sb_secret_rHJs88tXeSfpHaKPJdBPZg_7BeIqHbj"
PROGRAMA_ID = "95fc823b-0b37-4f42-8b5e-3f6f460b6382"

headers = {
    "apikey": SUPABASE_KEY,
    "Authorization": f"Bearer {SUPABASE_KEY}",
    "Content-Type": "application/json",
    "Prefer": "return=representation"
}

# Leer Excel
wb = load_workbook(r"C:\Users\Francisco Miranda\detalles-app\scripts\ACT RELEVANTES.xlsx", read_only=True)
ws = wb.active

# Construir mapa dia->frente por actividad
actividades_excel = {}

for i, row in enumerate(ws.iter_rows(values_only=True)):
    if i < 2:
        continue
    actividad = row[1]
    tipo_seq = row[3]
    imprimir = row[5]

    if imprimir != 'SI':
        continue
    if tipo_seq == 'FACHADA':
        continue

    dia_a_frente = {}
    for col_idx, val in enumerate(row[6:]):
        if val is not None:
            day_num = 100 + col_idx
            dia_a_frente[str(day_num)] = val

    actividades_excel[actividad] = dia_a_frente

print(f"Actividades leídas del Excel: {len(actividades_excel)}")

# Obtener actividades de Supabase
r = requests.get(
    f"{SUPABASE_URL}/rest/v1/programa_actividades?programa_id=eq.{PROGRAMA_ID}&select=id,actividad&limit=100",
    headers={
        "apikey": SUPABASE_KEY,
        "Authorization": f"Bearer {SUPABASE_KEY}",
    }
)
actividades_bd = r.json()
print(f"Actividades en BD: {len(actividades_bd)}")

if not isinstance(actividades_bd, list):
    print(f"Error inesperado: {actividades_bd}")
    exit()

# Actualizar cada actividad con su mapa
ok = 0
errores = []

for act_bd in actividades_bd:
    actividad_nombre = act_bd["actividad"]
    act_id = act_bd["id"]

    if actividad_nombre not in actividades_excel:
        errores.append(f"No encontrada en Excel: {actividad_nombre}")
        continue

    mapa = actividades_excel[actividad_nombre]

    r = requests.patch(
        f"{SUPABASE_URL}/rest/v1/programa_actividades?id=eq.{act_id}",
        headers=headers,
        json={"mapa_dias": mapa}
    )

    if r.status_code in (200, 204):
        ok += 1
        print(f"  ✓ {actividad_nombre}")
    else:
        errores.append(f"Error {actividad_nombre}: {r.text}")

print(f"\nActualizadas: {ok}")
if errores:
    print("Errores:")
    for e in errores:
        print(f"  - {e}")
